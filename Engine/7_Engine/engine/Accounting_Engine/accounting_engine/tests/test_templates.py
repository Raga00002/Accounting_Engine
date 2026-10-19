import io
from datetime import date

import openpyxl
import pytest

from engine import config, db, exporter, tally_export as tx, templates
from engine.models import Bill, LineItem, Party


@pytest.fixture()
def tpl_dir(tmp_path):
    """Activate a throw-away tenant so custom templates land under tmp_path."""
    from engine import tenant

    t = tenant.Tenant(id="test", root=tmp_path / "tenant")
    tok = tenant.activate(t)
    yield t.template_dir
    tenant.deactivate(tok)


def _xlsx(headers, sheet="Accounting Voucher"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_official_templates_are_bundled_and_valid():
    for kind in ("voucher", "masters"):
        inf = templates.info(kind)
        assert inf.source == "official" and inf.path and inf.path.exists(), kind
        assert set(templates.KINDS[kind]["required"]) <= set(inf.headers), kind
        assert inf.warnings == []
    assert templates.info("voucher").headers[:3] == ["Voucher Date", "Voucher Type Name", "Voucher Number"]
    assert templates.info("masters").headers[:2] == ["Name", "Group Name"]


def test_custom_template_validation(tpl_dir):
    with pytest.raises(ValueError, match="not an .xlsx"):
        templates.set_custom("voucher", b"hello")
    with pytest.raises(ValueError, match="mandatory"):
        templates.set_custom("voucher", _xlsx(["Voucher Date", "Ledger Name"]))
    with pytest.raises(ValueError, match="empty"):
        templates.set_custom("voucher", _xlsx([]))
    inf = templates.set_custom("voucher", _xlsx(["Voucher Date", "Voucher Type Name", "Ledger Name", "Ledger Amount", "My Column"]))
    assert inf.source == "custom" and templates.is_custom("voucher")
    assert any("Ledger Amount Dr/Cr" in w for w in inf.warnings)  # recommended column missing -> warned
    templates.clear_custom("voucher")
    assert templates.info("voucher").source == "official"


def test_custom_voucher_template_is_followed_exactly(tpl_dir):
    headers = ["Voucher Type Name", "Voucher Date", "Voucher Number", "Ledger Name", "Ledger Amount Dr/Cr",
               "Ledger Amount", "Item Name", "Billed Quantity", "Item Rate", "Item Amount", "Extra Col"]
    templates.set_custom("voucher", _xlsx(headers))
    party = Party(id=1, name="Das Store", group_name="Sundry Debtors")
    bill = Bill(id=1, party_id=1, bill_date=date(2026, 9, 26), bill_no="13", amount=7700, drcr="Dr",
                items=[LineItem(item="Rice (BPT)", qty=5, unit="Bag", rate=1540, amount=7700)])
    rows = tx.official_rows_for_bills([bill], {1: party}, tx.TallySettings())
    data = tx.official_xlsx(rows)
    ws = openpyxl.load_workbook(io.BytesIO(data))["Accounting Voucher"]
    assert [c.value for c in ws[1]][: len(headers)] == headers          # user's order kept
    assert ws.max_column == len(headers)                                 # nothing appended
    vals = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert vals[0][0] == "Sales" and vals[0][3] == "Das Store" and vals[0][4] == "Dr" and vals[0][5] == 7700
    # official template again -> our optional columns are appended after Tally's 14
    templates.clear_custom("voucher")
    ws = openpyxl.load_workbook(io.BytesIO(tx.official_xlsx(rows)))["Accounting Voucher"]
    hdr = [str(c.value).strip() for c in ws[1]]
    assert hdr[:14] == tx.OFFICIAL_BASE_HEADERS and "Bill Type of Ref" in hdr


def test_custom_masters_template_column_order(tpl_dir, tmp_db):
    headers = ["Group Name", "Name", "Address", "Pincode", "Bill - Name", "Bill - Date", "Bill - Amount", "Bill Amount - Dr/Cr"]
    templates.set_custom("masters", _xlsx(headers, sheet="Ledger"))
    p = db.upsert_party(Party(name="P. Roy", group_name="Sundry Debtors", address="Ward 3", pincode="786125"), tmp_db)
    db.insert_bill(Bill(party_id=p.id, bill_date=date(2026, 9, 26), bill_no="02", amount=100, drcr="Dr"), tmp_db)
    data, ids = exporter.export_bills(db_path=tmp_db)
    ws = openpyxl.load_workbook(io.BytesIO(data))["Ledger"]
    assert [c.value for c in ws[1]][: len(headers)] == headers
    row = [c.value for c in ws[2]]
    assert row[0] == "Sundry Debtors" and row[1] == "P. Roy" and row[2] == "Ward 3" and row[3] == "786125"
    assert row[4] == "02" and row[6] == 100 and row[7] == "Dr"


def test_custom_voucher_template_with_blank_header_column(tpl_dir):
    """A blank header cell in the middle must not shift the values into the wrong columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounting Voucher"
    for col, h in [(1, "Voucher Date"), (2, "Voucher Type Name"), (4, "Ledger Name"), (5, "Ledger Amount"), (6, "Ledger Amount Dr/Cr ")]:
        ws.cell(row=1, column=col, value=h)
    buf = io.BytesIO()
    wb.save(buf)
    templates.set_custom("voucher", buf.getvalue())
    party = Party(id=1, name="Das Store")
    bill = Bill(id=1, party_id=1, bill_date=date(2026, 9, 26), bill_no="13", amount=100, drcr="Dr")
    out = openpyxl.load_workbook(io.BytesIO(tx.official_xlsx(tx.official_rows_for_bills([bill], {1: party}, tx.TallySettings()))))
    ws = out["Accounting Voucher"]
    assert [c.value for c in ws[1]] == ["Voucher Date", "Voucher Type Name", None, "Ledger Name", "Ledger Amount", "Ledger Amount Dr/Cr "]
    r = [c.value for c in ws[2]]
    assert r[2] is None and r[3] == "Das Store" and r[4] == 100 and r[5] == "Dr"
    assert ws.max_column == 6


def test_custom_masters_template_never_writes_missing_columns_elsewhere(tpl_dir, tmp_db):
    """Opening balance / state / country are not in this template -> left out, not written into
    another column's position or appended after the template's columns."""
    headers = ["Group Name", "Name", "Address", "Pincode", "Bill - Name", "Bill - Date", "Bill - Amount", "Bill Amount - Dr/Cr"]
    templates.set_custom("masters", _xlsx(headers, sheet="Ledger"))
    p = db.upsert_party(Party(name="P. Roy", address="Ward 3", pincode="786125", opening_balance=500, opening_drcr="Dr",
                              state="Assam", country="India"), tmp_db)
    db.insert_bill(Bill(party_id=p.id, bill_date=date(2026, 9, 26), bill_no="02", amount=100, drcr="Dr"), tmp_db)
    data, _ = exporter.export_bills(db_path=tmp_db)
    ws = openpyxl.load_workbook(io.BytesIO(data))["Ledger"]
    assert ws.max_column == len(headers)
    assert [c.value for c in ws[2]][:4] == ["Sundry Debtors", "P. Roy", "Ward 3", "786125"]
