import io
from datetime import date, datetime

import openpyxl

from engine import config, db, exporter
from engine.importer import HEADERS, import_ledger_template
from engine.models import Bill, LineItem, Party


def test_export_matches_template_layout(tmp_db):
    abc = db.upsert_party(Party(name="ABC Company Limited", group_name="Sundry Debtors", opening_balance=10000,
                                opening_drcr="Dr", address="15 MG Road\nNear LIC building\nBangalore",
                                state="Karnataka", country="India"), tmp_db)
    titan = db.upsert_party(Party(name="Titan Corporation", group_name="Sundry Creditors", opening_balance=15000,
                                  opening_drcr="Cr", address="18, Raju Street\n1st main 3rd cross\nJP nagar 8th Phase",
                                  state="Karnataka", country="India", pincode="560078"), tmp_db)
    das = db.upsert_party(Party(name="Das Store", state="Assam", country="India"), tmp_db)
    for i, (d, amt) in enumerate([(date(2023, 3, 10), 1000), (date(2023, 3, 12), 2000), (date(2023, 3, 14), 7000)], 1):
        db.insert_bill(Bill(party_id=abc.id, bill_date=d, bill_no=f"ABC/{i}", amount=amt, drcr="Dr"), tmp_db)
    db.insert_bill(Bill(party_id=titan.id, bill_date=date(2023, 3, 20), bill_no="6542", amount=15000, drcr="Cr",
                        txn_type="purchase"), tmp_db)
    db.insert_bill(Bill(party_id=das.id, bill_date=date(2026, 9, 26), bill_no="13", amount=45125, drcr="Dr", source="image",
                        items=[LineItem(item="Rice (BPT)", qty=5, unit="Bag", rate=1540, amount=7700),
                               LineItem(item="Rice (Gobi)", qty=2, unit="Bag", weight_qtl=0.6, rate=4900, struck_out=True)]),
                   tmp_db)

    data, ids = exporter.export_bills(db_path=tmp_db, mark_exported=True)
    assert len(ids) == 5
    assert all(b.exported for b in db.list_bills(tmp_db))

    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Ledger", "Items"]
    ws = wb["Ledger"]
    assert [c.value for c in ws[1]][:12] == HEADERS
    assert ws["A1"].font.b and ws.column_dimensions["A"].width == 29.0  # template styling kept

    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=12)]
    # ABC block: 3 rows (3 bills / 3 address lines), same shape as the original template
    assert rows[0][:4] == ["ABC Company Limited", "Sundry Debtors", 10000, "Dr"]
    assert rows[0][4] == datetime(2023, 3, 10) and rows[0][5:8] == ["ABC/1", 1000, "Dr"]
    assert rows[0][8:12] == ["15 MG Road", "Karnataka", "India", None]
    assert rows[1][0] is None and rows[1][5] == "ABC/2" and rows[1][8] == "Near LIC building"
    assert rows[2][5] == "ABC/3" and rows[2][8] == "Bangalore"
    # Titan: 1 bill but 3 address lines -> 3 rows
    assert rows[3][:4] == ["Titan Corporation", "Sundry Creditors", 15000, "Cr"]
    assert rows[3][5:8] == ["6542", 15000, "Cr"] and rows[3][11] == "560078"
    assert rows[4][0] is None and rows[4][4] is None and rows[4][8] == "1st main 3rd cross"
    assert rows[5][8] == "JP nagar 8th Phase"
    # Das Store: no opening balance -> C/D blank
    assert rows[6][:4] == ["Das Store", "Sundry Debtors", None, None]
    assert rows[6][4] == datetime(2026, 9, 26) and rows[6][5:8] == ["13", 45125, "Dr"]
    assert len(rows) == 7
    assert ws["E2"].number_format == "dd/mmm/yyyy"

    items = wb["Items"]
    assert [c.value for c in items[1]] == exporter.ITEM_HEADERS
    irows = [[c.value for c in r] for r in items.iter_rows(min_row=2, max_row=items.max_row)]
    assert len(irows) == 2
    assert irows[0][:5] == ["Das Store", "13", datetime(2026, 9, 26), "sale", "Rice (BPT)"]
    assert irows[0][9] == 7700 and irows[0][12] == 45125
    assert irows[1][10] == "Yes" and irows[1][9] == 2940  # struck-out line kept for reference, flagged


def test_export_filters_and_options(tmp_db):
    import_ledger_template(config.TEMPLATE_PATH, tmp_db)  # 4 bills, all marked exported
    p = db.get_party_by_name("Das Store", tmp_db) or db.upsert_party(Party(name="Das Store"), tmp_db)
    db.insert_bill(Bill(party_id=p.id, bill_date=date(2026, 9, 26), bill_no="13", amount=45125, drcr="Dr"), tmp_db)

    data, ids = exporter.export_bills(db_path=tmp_db, only_unexported=True, include_opening_balance=False,
                                      include_address=False, include_items_sheet=False)
    assert len(ids) == 1
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Ledger"]
    ws = wb["Ledger"]
    assert ws.max_row == 2 and ws["A2"].value == "Das Store" and ws["C2"].value is None and ws["I2"].value is None

    data, ids = exporter.export_bills(db_path=tmp_db, date_from=date(2023, 3, 12), date_to=date(2023, 3, 20))
    assert len(ids) == 3  # ABC/2, ABC/3, 6542

    data, ids = exporter.export_bills(bill_ids=[], db_path=tmp_db)
    assert ids == [] and openpyxl.load_workbook(io.BytesIO(data))["Ledger"].max_row == 1


def test_round_trip_export_then_import(tmp_db, tmp_path):
    """What we export must be readable by our own importer (and therefore shaped like Tally's file)."""
    p = db.upsert_party(Party(name="P. Roy", address="Ward 3\nTinsukia", state="Assam", country="India",
                              pincode="786125"), tmp_db)
    db.insert_bill(Bill(party_id=p.id, bill_date=date(2026, 9, 26), bill_no="02", amount=79297, drcr="Dr"), tmp_db)
    out = tmp_path / "out.xlsx"
    exporter.export_bills(db_path=tmp_db, out_path=out)

    db2 = tmp_path / "second.db"
    db.init_db(db2)
    res = import_ledger_template(out, db2)
    assert res.parties_added == 1 and res.bills_added == 1
    q = db.get_party_by_name("P. Roy", db2)
    assert q.address == "Ward 3\nTinsukia" and q.pincode == "786125"
    assert db.list_bills(db2)[0].amount == 79297


def test_ui_dataframe_round_trip_is_stable():
    """items_to_df(df_to_items(df)) must not change the frame – the editor's identity depends on it."""
    import pyarrow as pa

    from ui.common import df_to_items, items_to_df

    items = [LineItem(item="Rice (BPT)", qty=5, unit="Bag", rate=1540, amount=7700),
             LineItem(item="Rice (Sri)", qty=2, unit="Bag", weight_qtl=0.52, rate=5550, struck_out=True)]
    df = items_to_df(items)
    assert list(df.dtypes.astype(str)) == ["object", "float64", "object", "float64", "float64", "float64", "bool", "object"]
    assert df_to_items(df) == items
    assert pa.Table.from_pandas(items_to_df(df_to_items(df))).equals(pa.Table.from_pandas(df))
    empty = items_to_df([])
    assert list(empty.dtypes.astype(str)) == list(df.dtypes.astype(str))
