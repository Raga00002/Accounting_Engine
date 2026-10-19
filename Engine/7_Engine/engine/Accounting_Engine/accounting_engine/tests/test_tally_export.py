import io
import zipfile
from datetime import date
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from engine import db, tally_export as tx
from engine.models import Bill, LineItem, Party


def _party():
    return Party(id=1, name="Das & Store", group_name="Sundry Debtors", address="Ward 3\nTinsukia",
                 state="Assam", country="India", pincode="786125", opening_balance=1500, opening_drcr="Dr")


def _bill(txn="sale"):
    return Bill(id=1, party_id=1, bill_date=date(2026, 9, 26), bill_no="13", amount=16346, drcr="Dr" if txn == "sale" else "Cr",
                txn_type=txn, total_written=16346, items=[
                    LineItem(item="Rice (BPT)", qty=5, unit="Bag", rate=1540, amount=7700),
                    LineItem(item="Rice (Sri)", qty=2, unit="Bag", weight_qtl=0.52, rate=5550, amount=2886),
                    LineItem(item="Rice (Gobi)", qty=2, unit="Bag", weight_qtl=0.6, rate=4900, struck_out=True),
                    LineItem(item="Mung Dal", qty=2, unit="Bag", weight_qtl=0.60, rate=9600),  # amount computed
                ])


def test_inventory_lines_weight_and_package_pricing():
    s = tx.TallySettings()
    lines = tx.inventory_lines(_bill(), s)
    assert [(l.item, l.qty, l.unit, l.rate, l.amount) for l in lines] == [
        ("Rice (BPT)", 5.0, "Bag", 1540.0, 7700.0),
        ("Rice (Sri)", 0.52, "Qtl", 5550.0, 2886.0),
        ("Mung Dal", 0.6, "Qtl", 9600.0, 5760.0),
    ]
    # unit falls back to the items master, then to the default unit; rate derived when missing
    b = Bill(id=2, party_id=1, bill_date=date(2026, 1, 1), bill_no="x", amount=300, drcr="Dr",
             items=[LineItem(item="Oil", qty=3, amount=300), LineItem(item="Salt", amount=50)])
    lines = tx.inventory_lines(b, s, item_units={"oil": "Tin"})
    assert (lines[0].unit, lines[0].rate) == ("Tin", 100.0)
    assert (lines[1].qty, lines[1].unit, lines[1].rate) == (1.0, "Nos", 50.0)


def _voucher(xml: str):
    root = ET.fromstring(xml.encode("utf-8"))
    return root.find(".//VOUCHER")


def test_sales_voucher_xml_is_balanced_and_signed_correctly():
    s = tx.TallySettings(company="Wings Industries", godown="Main Location")
    xml = tx.vouchers_xml([_bill()], {1: _party()}, s)
    assert tx.validate_xml(xml.encode()) == 1
    assert "<SVCURRENTCOMPANY>Wings Industries</SVCURRENTCOMPANY>" in xml
    v = _voucher(xml)
    assert v.get("VCHTYPE") == "Sales" and v.findtext("DATE") == "20260926" and v.findtext("VOUCHERNUMBER") == "13"
    assert v.findtext("PARTYLEDGERNAME") == "Das & Store"  # escaped in XML, unescaped on parse
    led = v.findall("LEDGERENTRIES.LIST")
    assert led[0].findtext("LEDGERNAME") == "Das & Store"
    assert led[0].findtext("ISDEEMEDPOSITIVE") == "Yes" and led[0].findtext("AMOUNT") == "-16346.00"
    assert led[0].find("BILLALLOCATIONS.LIST").findtext("BILLTYPE") == "New Ref"
    inv = v.findall("ALLINVENTORYENTRIES.LIST")
    assert [e.findtext("STOCKITEMNAME") for e in inv] == ["Rice (BPT)", "Rice (Sri)", "Mung Dal"]
    assert inv[1].findtext("RATE") == "5550.00/Qtl" and inv[1].findtext("ACTUALQTY") == "0.52 Qtl"
    assert inv[0].find("BATCHALLOCATIONS.LIST").findtext("GODOWNNAME") == "Main Location"
    # debit + credits = 0
    total = sum(float(e.findtext("AMOUNT")) for e in led) + sum(float(e.findtext("AMOUNT")) for e in inv)
    assert total == pytest.approx(0)
    assert all(e.find("ACCOUNTINGALLOCATIONS.LIST").findtext("LEDGERNAME") == "Sales" for e in inv)


def test_purchase_voucher_and_accounting_only_mode():
    s = tx.TallySettings(inventory=False)
    xml = tx.vouchers_xml([_bill("purchase")], {1: _party()}, s)
    v = _voucher(xml)
    assert v.get("VCHTYPE") == "Purchase" and v.findtext("ISINVOICE") == "No"
    assert v.get("OBJVIEW") == "Accounting Voucher View" and v.findtext("VCHENTRYMODE") == "As Voucher"
    # voucher mode = ALLLEDGERENTRIES.LIST (what Tally itself exports for Accounting Voucher View)
    assert v.findall("LEDGERENTRIES.LIST") == []
    led = v.findall("ALLLEDGERENTRIES.LIST")
    assert led[0].findtext("AMOUNT") == "16346.00" and led[0].findtext("ISDEEMEDPOSITIVE") == "No"  # party credited
    assert led[0].find("BILLALLOCATIONS.LIST").findtext("AMOUNT") == "16346.00"
    assert led[1].findtext("LEDGERNAME") == "Purchase" and led[1].findtext("AMOUNT") == "-16346.00"
    assert "Rice (BPT) 5 Bag @ 1540.00" in v.findtext("NARRATION")  # items kept in narration
    assert v.findall("ALLINVENTORYENTRIES.LIST") == []
    assert v.find("BASICBUYERNAME") is None and v.findtext("BASICBASEPARTYNAME") == "Das & Store"


def test_purchase_invoice_with_inventory_is_balanced():
    s = tx.TallySettings()
    v = _voucher(tx.vouchers_xml([_bill("purchase")], {1: _party()}, s))
    led = v.findall("LEDGERENTRIES.LIST")
    assert led[0].findtext("AMOUNT") == "16346.00" and led[0].findtext("ISDEEMEDPOSITIVE") == "No"
    inv = v.findall("ALLINVENTORYENTRIES.LIST")
    assert all(e.findtext("ISDEEMEDPOSITIVE") == "Yes" and float(e.findtext("AMOUNT")) < 0 for e in inv)
    assert all(e.find("ACCOUNTINGALLOCATIONS.LIST").findtext("LEDGERNAME") == "Purchase" for e in inv)
    assert all(e.find("ACCOUNTINGALLOCATIONS.LIST").findtext("ISDEEMEDPOSITIVE") == "Yes" for e in inv)
    total = sum(float(e.findtext("AMOUNT")) for e in led) + sum(float(e.findtext("AMOUNT")) for e in inv)
    assert total == pytest.approx(0)


def test_number_formatting_and_escaping():
    assert tx._amt(-0.0) == "0.00" and tx._amt(2.675) in ("2.67", "2.68") and tx._amt(0.1 + 0.2) == "0.30"
    assert tx._fmt_qty(13.10) == "13.1" and tx._fmt_qty(0.1 + 0.2) == "0.3" and tx._fmt_qty(5.0) == "5"
    assert tx._x('A & B <"C">\x0bD') == "A &amp; B &lt;&quot;C&quot;&gt;D"
    # zero / non-finite quantities never divide or reach the XML; float noise is rounded away
    b = Bill(id=3, party_id=1, bill_date=date(2026, 1, 1), bill_no="n", amount=10, drcr="Dr", items=[
        LineItem(item="Zero", qty=0, rate=5, amount=10),  # qty 0 -> treated as 1 x 10
        LineItem(item="Noisy", qty=0.1 + 0.2, unit="Kg", amount=0.3),
        LineItem(item="Inf", qty=1, amount=float("inf")),
    ])
    lines = tx.inventory_lines(b, tx.TallySettings())
    assert [(l.item, l.qty, l.rate) for l in lines] == [("Zero", 1.0, 10.0), ("Noisy", 0.3, 1.0)]
    xml = tx.vouchers_xml([b], {1: _party()}, tx.TallySettings())
    assert "0.3 Kg" in xml and tx.validate_xml(xml.encode()) == 1
    # non-ASCII survives the UTF-8 round trip
    p = Party(id=1, name="Śrī Kṛṣṇa Traders", group_name="")
    xml = tx.masters_xml([p], [], tx.TallySettings()).encode("utf-8")
    root = ET.fromstring(xml)
    assert root.find(".//LEDGER").get("NAME") == "Śrī Kṛṣṇa Traders"
    assert root.findtext(".//LEDGER/PARENT") == "Sundry Debtors"  # blank group never produces an empty PARENT


def test_balancing_difference_goes_to_sales_ledger():
    b = _bill()
    b.amount = 16350  # booked total differs from lines by 4
    xml = tx.vouchers_xml([b], {1: _party()}, tx.TallySettings())
    v = _voucher(xml)
    led = v.findall("LEDGERENTRIES.LIST")
    assert led[-1].findtext("LEDGERNAME") == "Sales" and led[-1].findtext("AMOUNT") == "4.00"
    amounts = [float(e.findtext("AMOUNT")) for e in led] + [float(e.findtext("AMOUNT")) for e in v.findall("ALLINVENTORYENTRIES.LIST")]
    assert sum(amounts) == pytest.approx(0)


def test_negative_balancing_difference_flips_side():
    # booked total below the line total: Sales must be *debited* (Yes / negative), Purchase credited
    b = _bill()
    b.amount = 16340  # lines sum to 16346 -> diff -6
    v = _voucher(tx.vouchers_xml([b], {1: _party()}, tx.TallySettings()))
    led = v.findall("LEDGERENTRIES.LIST")
    assert led[-1].findtext("LEDGERNAME") == "Sales"
    assert led[-1].findtext("ISDEEMEDPOSITIVE") == "Yes" and led[-1].findtext("AMOUNT") == "-6.00"
    amounts = [float(e.findtext("AMOUNT")) for e in led] + [float(e.findtext("AMOUNT")) for e in v.findall("ALLINVENTORYENTRIES.LIST")]
    assert sum(amounts) == pytest.approx(0)
    bp = _bill("purchase")
    bp.amount = 16340
    v = _voucher(tx.vouchers_xml([bp], {1: _party()}, tx.TallySettings()))
    led = v.findall("LEDGERENTRIES.LIST")
    assert led[-1].findtext("ISDEEMEDPOSITIVE") == "No" and led[-1].findtext("AMOUNT") == "6.00"
    # every ledger / inventory entry in every mode has a flag consistent with its sign
    for bill in (b, bp, _bill(), _bill("purchase")):
        for inv in (True, False):
            v = _voucher(tx.vouchers_xml([bill], {1: _party()}, tx.TallySettings(inventory=inv)))
            for e in v.iter():
                if e.tag in ("LEDGERENTRIES.LIST", "ALLLEDGERENTRIES.LIST", "ALLINVENTORYENTRIES.LIST", "ACCOUNTINGALLOCATIONS.LIST"):
                    amt = float(e.findtext("AMOUNT"))
                    assert (e.findtext("ISDEEMEDPOSITIVE") == "Yes") == (amt < 0), (e.tag, amt)
    # Excel / CSV rows: no negative amounts, the difference row is on the opposite side, Dr == Cr
    rows = tx.voucher_rows([b], {1: _party()}, tx.TallySettings())
    assert rows[-1][4:7] == ["Sales", 6, "Dr"] and all(r[5] > 0 for r in rows)
    assert sum(r[5] for r in rows if r[6] == "Dr") == pytest.approx(sum(r[5] for r in rows if r[6] == "Cr"))


def test_unit_fallbacks_and_tiny_quantities():
    s = tx.TallySettings()
    b = Bill(id=4, party_id=1, bill_date=date(2026, 1, 1), bill_no="u", amount=100, drcr="Dr", items=[
        LineItem(item="Blank unit", qty=2, unit="  ", rate=50, amount=100),  # whitespace unit -> default unit
        LineItem(item="Tiny", qty=0.00001, amount=1),  # rounds to 0 at 4 dp -> skipped, never "0 Nos"
    ])
    lines = tx.inventory_lines(b, s)
    assert [(l.item, l.unit) for l in lines] == [("Blank unit", "Nos")]
    xml = tx.vouchers_xml([b], {1: _party()}, s)
    assert "2 Nos" in xml and "0 Nos" not in xml and "/Nos" in xml
    # a stock item never gets an empty BASEUNITS, even with a blank default unit
    m = ET.fromstring(tx.masters_xml([], [("Oil", None, None)], tx.TallySettings(default_unit="")).encode())
    assert m.findtext(".//STOCKITEM/BASEUNITS") == "Nos" and [u.get("NAME") for u in m.findall(".//UNIT")] == ["Nos"]
    assert m.findtext(".//UNIT/DECIMALPLACES") == "4"  # same precision as the exported quantities


def test_masters_xml():
    s = tx.TallySettings(stock_group="Grains")
    xml = tx.masters_xml([_party()], [("Rice (BPT)", "Bag", None), ("Mung Dal", "Qtl", "Pulses")], s)
    root = ET.fromstring(xml.encode())
    assert root.findtext(".//IMPORTDUPS") == "@@DUPIGNORECOMBINE"
    led = root.find(".//LEDGER")
    assert led.get("NAME") == "Das & Store" and led.findtext("PARENT") == "Sundry Debtors"
    assert [a.text for a in led.findall("ADDRESS.LIST/ADDRESS")] == ["Ward 3", "Tinsukia"]
    assert led.findtext("OPENINGBALANCE") == "-1500.00"  # Dr opening is negative
    assert sorted(u.get("NAME") for u in root.findall(".//UNIT")) == ["Bag", "Qtl"]
    assert sorted(g.get("NAME") for g in root.findall(".//STOCKGROUP")) == ["Grains", "Pulses"]
    items = {i.get("NAME"): i for i in root.findall(".//STOCKITEM")}
    assert items["Rice (BPT)"].findtext("PARENT") == "Grains" and items["Mung Dal"].findtext("PARENT") == "Pulses"
    assert items["Rice (BPT)"].findtext("BASEUNITS") == "Bag"


def test_excel_and_csv_rows():
    s = tx.TallySettings()
    data = tx.vouchers_xlsx([_bill()], {1: _party()}, s)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Vouchers", "ReadMe"]
    ws = wb["Vouchers"]
    assert [c.value for c in ws[1]] == tx.VOUCHER_HEADERS
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 4  # party + 3 items
    assert rows[0][4:9] == ["Das & Store", 16346, "Dr", "13", "New Ref"]
    assert rows[1][4:7] == ["Sales", 7700, "Cr"] and rows[1][9:14] == ["Rice (BPT)", 5, "Bag", 1540, 7700]
    assert rows[0][0].strftime("%Y-%m-%d") == "2026-09-26"
    # ledger debits == credits
    assert sum(r[5] for r in rows if r[6] == "Dr") == sum(r[5] for r in rows if r[6] == "Cr")
    csv_text = tx.vouchers_csv([_bill()], {1: _party()}, s).decode("utf-8-sig")
    lines = csv_text.splitlines()
    assert lines[0] == ",".join(tx.VOUCHER_HEADERS)
    assert lines[1].startswith("26-09-2026,Sales,13,Das & Store,Das & Store,16346.0,Dr,13,New Ref")


def test_bundle_from_database(tmp_db):
    p = db.upsert_party(Party(name="P. Roy", group_name="Sundry Debtors"), tmp_db)
    db.upsert_item("Mung Dal", default_unit="Pkt", group_name="Pulses", db_path=tmp_db)
    b = db.insert_bill(Bill(party_id=p.id, bill_date=date(2026, 9, 26), bill_no="02", amount=42975, drcr="Dr",
                            items=[LineItem(item="Mung Dal", qty=15, weight_qtl=4.5, rate=9550)]), tmp_db)
    s = tx.TallySettings(company="Wings", cash_ledger=" Cash-in-Hand ", contra_vch_type="Contra", receipt_vch_type="Receipt")
    s.save(tmp_db)
    loaded = tx.TallySettings.load(tmp_db)
    assert loaded.company == "Wings" and loaded.cash_ledger == "Cash-in-Hand"  # full round trip incl. the bank fields
    assert loaded == tx.TallySettings(company="Wings", cash_ledger="Cash-in-Hand")
    tx.TallySettings(cash_ledger="").save(tmp_db)
    assert tx.TallySettings.load(tmp_db).cash_ledger == "Cash"  # blank -> default, never an empty ledger name
    s.save(tmp_db)
    bundle = tx.build_bundle([b], list(tx.FORMATS), db_path=tmp_db, stem="T")
    assert set(bundle.files) == {"T.xml", "T_masters.xml", "T.xlsx", "T_TallyTemplate.xlsx", "T.csv"} and bundle.bill_ids == [b.id]
    assert tx.validate_xml(bundle.files["T.xml"]) == 1
    m = ET.fromstring(bundle.files["T_masters.xml"])
    item = m.find(".//STOCKITEM")
    assert item.get("NAME") == "Mung Dal" and item.findtext("BASEUNITS") == "Qtl" and item.findtext("PARENT") == "Pulses"
    # items master says Pkt, weight-priced line is exported in Qtl -> surfaced, not silent
    assert any("items master unit is 'Pkt'" in w and "Qtl" in w for w in bundle.warnings), bundle.warnings
    # the Excel quantity column is numeric with a format that never prints "15."
    ws = openpyxl.load_workbook(io.BytesIO(bundle.files["T.xlsx"]))["Vouchers"]
    assert ws.cell(row=3, column=11).value == 4.5 and ws.cell(row=3, column=11).number_format == "General"
    z = zipfile.ZipFile(io.BytesIO(bundle.zip_bytes()))
    assert sorted(z.namelist()) == ["T.csv", "T.xlsx", "T.xml", "T_TallyTemplate.xlsx", "T_masters.xml"]
    # a bill whose party vanished is skipped with a warning, not a crash
    orphan = Bill(id=99, party_id=999, bill_date=date(2026, 1, 1), bill_no="z", amount=1, drcr="Dr")
    bundle = tx.build_bundle([orphan], ["voucher_xml"], db_path=tmp_db)
    assert bundle.files == {} and bundle.warnings and bundle.bill_ids == []


def test_unit_warnings():
    assert tx.unit_warnings({"oil": {"Tin"}}, {"oil": "Tin"}) == []
    assert tx.unit_warnings({"oil": {"tin"}}, {"oil": "Tin"}) == []  # case-insensitive
    w = tx.unit_warnings({"oil": {"Tin", "Ltr"}, "salt": {"Kg"}}, {"oil": "Tin"})
    assert len(w) == 2 and "several units" in w[0] and "items master unit is 'Tin'" in w[1] and "Ltr" in w[1]


def test_official_template_xlsx_for_bills():
    s = tx.TallySettings()
    data = tx.official_xlsx(tx.official_rows_for_bills([_bill()], {1: _party()}, s))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Accounting Voucher" and "Accounting Voucher (Read Me)" in wb.sheetnames
    ws = wb["Accounting Voucher"]
    headers = [str(c.value) for c in ws[1] if c.value is not None]
    # Tally's own 14 headers first (exact spelling, incl. the trailing space of 'Change Mode ' in the sample file)
    assert [h.strip() for h in headers[:14]] == tx.OFFICIAL_BASE_HEADERS
    assert headers[13] == "Change Mode "  # untouched template cell
    for h in tx.OFFICIAL_EXTRA_HEADERS:
        assert h in headers
    col = {h.strip(): i for i, h in enumerate(headers)}
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=ws.max_row)]
    assert len(rows) == 4
    r0 = rows[0]
    assert r0[col["Ledger Name"]] == "Das & Store" and r0[col["Ledger Amount"]] == 16346 and r0[col["Ledger Amount Dr/Cr"]] == "Dr"
    assert r0[col["Bill Type of Ref"]] == "New Ref" and r0[col["Bill Name"]] == "13" and r0[col["Change Mode"]] == "Item Invoice"
    assert r0[col["Buyer/Supplier - Address"]] == "Ward 3, Tinsukia" and r0[col["Buyer/Supplier - Pincode"]] == "786125"
    assert r0[col["Voucher Date"]].strftime("%Y-%m-%d") == "2026-09-26" and r0[col["Voucher Type Name"]] == "Sales"
    r1 = rows[1]
    assert (r1[col["Ledger Name"]], r1[col["Ledger Amount"]], r1[col["Ledger Amount Dr/Cr"]]) == ("Sales", 7700, "Cr")
    assert (r1[col["Item Name"]], r1[col["Billed Quantity"]], r1[col["Item Rate"]], r1[col["Item Rate per"]], r1[col["Item Amount"]]) == ("Rice (BPT)", 5, 1540, "Bag", 7700)
    assert r1[col["Quantity UOM"]] == "Bag"
    assert rows[2][col["Item Rate per"]] == "Qtl" and rows[2][col["Billed Quantity"]] == 0.52
    # Dr == Cr per voucher
    dr = sum(r[col["Ledger Amount"]] for r in rows if r[col["Ledger Amount Dr/Cr"]] == "Dr")
    cr = sum(r[col["Ledger Amount"]] for r in rows if r[col["Ledger Amount Dr/Cr"]] == "Cr")
    assert dr == pytest.approx(cr)


def test_simple_vouchers_contra_and_receipt():
    s = tx.TallySettings(company="Wings")
    contra = tx.SimpleVoucher(
        date=date(2026, 9, 19), vch_type="Contra", number="DEP/001", narration="Cash deposit PNB Nagaon",
        entries=[tx.SimpleEntry("PNB A/c 2218", 140000, "Dr", bank=tx.BankAlloc("Cash", "", None, "Punjab National Bank", "Nagaon")),
                 tx.SimpleEntry("Cash", 140000, "Cr")])
    receipt = tx.SimpleVoucher(
        date=date(2026, 9, 25), vch_type="Receipt", number="RCT/002",
        entries=[tx.SimpleEntry("PNB A/c 2218", 350000, "Dr", bank=tx.BankAlloc("Cash", bank_name="Punjab National Bank")),
                 tx.SimpleEntry("Das Store", 350000, "Cr", is_party=True, bill_type="On Account")])
    assert contra.balanced() and receipt.balanced()
    xml = tx.simple_vouchers_xml([contra, receipt], s)
    assert tx.validate_xml(xml.encode()) == 2
    root = ET.fromstring(xml.encode())
    v1, v2 = root.findall(".//VOUCHER")
    assert v1.get("VCHTYPE") == "Contra" and v1.findtext("DATE") == "20260919"
    e = v1.findall("ALLLEDGERENTRIES.LIST")
    assert e[0].findtext("LEDGERNAME") == "PNB A/c 2218" and e[0].findtext("AMOUNT") == "-140000.00" and e[0].findtext("ISDEEMEDPOSITIVE") == "Yes"
    ba = e[0].find("BANKALLOCATIONS.LIST")
    assert ba.findtext("TRANSACTIONTYPE") == "Cash" and ba.findtext("BANKNAME") == "Punjab National Bank" and ba.findtext("AMOUNT") == "-140000.00"
    assert e[1].findtext("LEDGERNAME") == "Cash" and e[1].findtext("AMOUNT") == "140000.00" and e[1].find("BANKALLOCATIONS.LIST") is None
    assert v2.findtext("PARTYLEDGERNAME") == "Das Store"
    pe = v2.findall("ALLLEDGERENTRIES.LIST")[1]
    ba = pe.find("BILLALLOCATIONS.LIST")  # bill-wise party: explicit On Account allocation, no bill name (as Tally exports it)
    assert pe.findtext("ISPARTYLEDGER") == "Yes" and ba.findtext("BILLTYPE") == "On Account" and ba.find("NAME") is None
    assert ba.findtext("AMOUNT") == "350000.00"  # same sign as the ledger entry (Cr = positive)
    agst = tx.SimpleVoucher(date=date(2026, 9, 25), vch_type="Receipt", number="RCT/003",
                            entries=[tx.SimpleEntry("PNB A/c 2218", 100, "Dr", bank=tx.BankAlloc()),
                                     tx.SimpleEntry("Das Store", 100, "Cr", is_party=True, bill_type="Agst Ref", bill_name="13")])
    a = ET.fromstring(tx.simple_vouchers_xml([agst], s).encode()).find(".//ALLLEDGERENTRIES.LIST[2]/BILLALLOCATIONS.LIST")
    assert a.findtext("NAME") == "13" and a.findtext("BILLTYPE") == "Agst Ref" and a.findtext("AMOUNT") == "100.00"
    for v in (v1, v2):
        assert sum(float(x.findtext("AMOUNT")) for x in v.findall("ALLLEDGERENTRIES.LIST")) == pytest.approx(0)

    bundle = tx.build_simple_bundle([contra, receipt], list(tx.FORMATS), s, stem="B", bank_ledgers=["PNB A/c 2218"],
                                    parties=[Party(id=1, name="Das Store")])
    assert set(bundle.files) == {"B.xml", "B_masters.xml", "B_TallyTemplate.xlsx", "B.xlsx", "B.csv"} and not bundle.warnings
    m = ET.fromstring(bundle.files["B_masters.xml"])
    names = {l.get("NAME"): l.findtext("PARENT") for l in m.findall(".//LEDGER")}
    assert names == {"PNB A/c 2218": "Bank Accounts", "Das Store": "Sundry Debtors"}
    wb = openpyxl.load_workbook(io.BytesIO(bundle.files["B_TallyTemplate.xlsx"]))
    ws = wb["Accounting Voucher"]
    headers = [str(c.value).strip() for c in ws[1] if c.value is not None]
    col = {h: i for i, h in enumerate(headers)}
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=ws.max_row)]
    assert len(rows) == 4
    assert rows[0][col["Voucher Type Name"]] == "Contra" and rows[0][col["Change Mode"]] == "As Voucher"
    assert rows[0][col["Bank Allocations - Transaction Type"]] == "Cash" and rows[0][col["Bank Allocations - Bank Name"]] == "Punjab National Bank"
    assert rows[0][col["Bank Allocations - Amount"]] == 140000 and rows[1][col["Ledger Name"]] == "Cash"
    assert rows[3][col["Bill Type of Ref"]] == "On Account" and rows[3][col["Ledger Amount Dr/Cr"]] == "Cr"
    csv_lines = bundle.files["B.csv"].decode("utf-8-sig").splitlines()
    assert csv_lines[1].startswith("19-09-2026,Contra,DEP/001,,PNB A/c 2218,140000,Dr")
    # an unbalanced voucher is flagged, never silently exported
    bad = tx.SimpleVoucher(date=date(2026, 1, 1), vch_type="Contra", number="x",
                           entries=[tx.SimpleEntry("Bank", 10, "Dr"), tx.SimpleEntry("Cash", 9, "Cr")])
    assert tx.build_simple_bundle([bad], ["voucher_xml"], s, stem="x").warnings
