"""Export saved bills as Tally-importable vouchers.

Formats
-------
* **Voucher XML**  – Tally's own import format (Gateway of Tally → Import → Transactions, XML).
  One <VOUCHER> per bill: party ledger entry with a bill-wise "New Ref", the sales/purchase
  ledger, and (optionally) one <ALLINVENTORYENTRIES.LIST> per line item.
* **Masters XML**  – <LEDGER>, <UNIT> and <STOCKITEM> objects for the parties/items used by the
  exported bills, so a voucher import never fails on a missing master.
* **Voucher Excel** – one row per entry (party row, then one row per stock item with its sales
  ledger allocation) using TallyPrime's mandatory field names ("Voucher Type Name", "Voucher Date",
  "Ledger Name", "Ledger Amount"…). Import → Transactions → Excel; map the columns once, Tally
  remembers the mapping.
* **Voucher CSV**  – the same rows as CSV (for spreadsheets / other software; Tally itself wants
  the .xlsx).

Sign conventions in Tally XML: a debit is a NEGATIVE amount with ISDEEMEDPOSITIVE=Yes, a credit a
POSITIVE amount with ISDEEMEDPOSITIVE=No. Dates are YYYYMMDD. Quantities are "<n> <unit>" and
rates "<amount>/<unit>".
"""
from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from typing import Iterable, Optional, Sequence
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, db
from .models import Bill, LineItem, Party

# --------------------------------------------------------------------------- #
# settings (persisted in the settings table; edited on the Settings page)
# --------------------------------------------------------------------------- #
SETTING_KEYS = {
    "tally_company": "",  # blank = whichever company is open in Tally
    "tally_sales_ledger": "Sales",
    "tally_purchase_ledger": "Purchase",
    "tally_sales_vch_type": "Sales",
    "tally_purchase_vch_type": "Purchase",
    "tally_inventory": "yes",  # yes = stock items on the voucher; no = accounting-only voucher
    "tally_godown": "",  # blank = Tally's default (Main Location)
    "tally_default_unit": "Nos",  # used when a line has no unit
    "tally_weight_unit": "Qtl",  # unit used when a line is priced by weight
    "tally_stock_group": "",  # parent stock group for new stock items in masters XML (blank = Primary)
    "tally_cash_ledger": "Cash",  # counter ledger for cash deposits (Contra vouchers)
    "tally_contra_vch_type": "Contra",
    "tally_receipt_vch_type": "Receipt",
}


@dataclass
class TallySettings:
    company: str = ""
    sales_ledger: str = "Sales"
    purchase_ledger: str = "Purchase"
    sales_vch_type: str = "Sales"
    purchase_vch_type: str = "Purchase"
    inventory: bool = True
    godown: str = ""
    default_unit: str = "Nos"
    weight_unit: str = "Qtl"
    stock_group: str = ""
    cash_ledger: str = "Cash"
    contra_vch_type: str = "Contra"
    receipt_vch_type: str = "Receipt"

    @classmethod
    def load(cls, db_path=None) -> "TallySettings":
        g = lambda k: (db.get_setting(k, SETTING_KEYS[k], db_path) or SETTING_KEYS[k]).strip()  # noqa: E731
        return cls(
            company=g("tally_company"),
            sales_ledger=g("tally_sales_ledger") or "Sales",
            purchase_ledger=g("tally_purchase_ledger") or "Purchase",
            sales_vch_type=g("tally_sales_vch_type") or "Sales",
            purchase_vch_type=g("tally_purchase_vch_type") or "Purchase",
            inventory=g("tally_inventory").lower() in ("yes", "true", "1"),
            godown=g("tally_godown"),
            default_unit=g("tally_default_unit") or "Nos",
            weight_unit=g("tally_weight_unit") or "Qtl",
            stock_group=g("tally_stock_group"),
            cash_ledger=g("tally_cash_ledger") or "Cash",
            contra_vch_type=g("tally_contra_vch_type") or "Contra",
            receipt_vch_type=g("tally_receipt_vch_type") or "Receipt",
        )

    def save(self, db_path=None) -> None:
        vals = {
            "tally_company": self.company,
            "tally_sales_ledger": self.sales_ledger,
            "tally_purchase_ledger": self.purchase_ledger,
            "tally_sales_vch_type": self.sales_vch_type,
            "tally_purchase_vch_type": self.purchase_vch_type,
            "tally_inventory": "yes" if self.inventory else "no",
            "tally_godown": self.godown,
            "tally_default_unit": self.default_unit,
            "tally_weight_unit": self.weight_unit,
            "tally_stock_group": self.stock_group,
            "tally_cash_ledger": self.cash_ledger,
            "tally_contra_vch_type": self.contra_vch_type,
            "tally_receipt_vch_type": self.receipt_vch_type,
        }
        for k, v in vals.items():
            db.set_setting(k, (v or "").strip(), db_path)


# --------------------------------------------------------------------------- #
# per-line quantity / rate resolution
# --------------------------------------------------------------------------- #
@dataclass
class InvLine:
    item: str
    qty: float
    unit: str
    rate: float
    amount: float
    note: str = ""


def inventory_lines(bill: Bill, s: TallySettings, item_units: Optional[dict[str, str]] = None) -> list[InvLine]:
    """Turn bill lines into Tally inventory lines.

    Weight-priced lines (".52 x 5550") become 0.52 <weight_unit> @ 5550/<weight_unit>; package-priced
    lines become <qty> <unit> @ rate. A line without rate gets rate = amount / qty. Struck-out lines
    and lines without a determinable amount are skipped.
    """
    item_units = item_units or {}
    out: list[InvLine] = []
    for it in bill.items:
        if it.struck_out:
            continue
        amount = it.effective_amount
        if amount is None or not math.isfinite(amount):
            continue
        amount = round(float(amount), 2)
        if it.weight_qtl is not None and it.weight_qtl > 0 and it.rate is not None:
            qty, unit, rate = float(it.weight_qtl), s.weight_unit, float(it.rate)
        else:
            qty = float(it.qty) if it.qty else 1.0
            # "  " is truthy: strip each candidate, or a blank unit gives "5 " / "150.00/" in the XML
            unit = (it.unit or "").strip() or item_units.get(it.item.lower(), "").strip() or s.default_unit.strip() or "Nos"
            rate = float(it.rate) if it.rate is not None and it.qty else amount / qty
        if not (math.isfinite(qty) and math.isfinite(rate)):
            continue
        # 4 dp = Tally's maximum unit precision; also removes 0.1+0.2-style float noise.
        # Round *before* the zero test: 0.00001 would otherwise pass it and reach the XML as "0 Qtl".
        qty, rate = round(qty, 4), round(rate, 4)
        if qty == 0:
            continue
        note = ""
        if it.qty and it.weight_qtl is not None and it.weight_qtl > 0:
            note = f"{_fmt_qty(it.qty)} {it.unit or 'Bag'}"
        out.append(InvLine(item=it.item.strip(), qty=qty, unit=unit, rate=rate, amount=amount, note=note))
    return out


def _fmt_qty(q: float) -> str:
    return f"{q:.4f}".rstrip("0").rstrip(".") if q != int(q) else str(int(q))


def _amt(v: float) -> str:
    v = round(float(v), 2)
    return f"{v + 0.0 if v != 0 else 0.0:.2f}"  # never "-0.00"


def _tally_date(d) -> str:
    return d.strftime("%Y%m%d")


def _vch(bill: Bill, s: TallySettings) -> tuple[str, str, str]:
    """(voucher type, sales/purchase ledger, party dr/cr sign) for the bill."""
    if bill.txn_type == "purchase":
        return s.purchase_vch_type, s.purchase_ledger, "Cr"
    return s.sales_vch_type, s.sales_ledger, "Dr"


def _narration(bill: Bill, party: Party, lines: list[InvLine], include_items: bool) -> str:
    bits = [f"Bill {bill.bill_no} – {party.name}"]
    if include_items and lines:
        bits.append("; ".join(f"{ln.item} {_fmt_qty(ln.qty)} {ln.unit} @ {_amt(ln.rate)}" for ln in lines))
    if bill.total_written is not None and abs(bill.total_written - bill.amount) > 0.5:
        bits.append(f"(written total {_amt(bill.total_written)})")
    return " · ".join(bits)


# --------------------------------------------------------------------------- #
# XML
# --------------------------------------------------------------------------- #
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _x(v) -> str:
    # control characters are illegal in XML 1.0 (a stray \x0b from an OCR'd name would make the
    # whole file unreadable for Tally); &, <, >, " are escaped
    return escape(_CTRL.sub("", str(v)), {'"': "&quot;"})


def _envelope(messages: Iterable[str], report: str, company: str, dups: str = "") -> str:
    sv = f"<SVCURRENTCOMPANY>{_x(company)}</SVCURRENTCOMPANY>" if company else ""
    # Masters: ignore names that already exist in Tally instead of overwriting them.
    sv += f"<IMPORTDUPS>{dups}</IMPORTDUPS>" if dups else ""
    body = "\n".join(messages)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<ENVELOPE>\n<HEADER><TALLYREQUEST>Import Data</TALLYREQUEST></HEADER>\n<BODY>\n<IMPORTDATA>\n"
        f"<REQUESTDESC><REPORTNAME>{report}</REPORTNAME><STATICVARIABLES>{sv}"
        "</STATICVARIABLES></REQUESTDESC>\n"
        f"<REQUESTDATA>\n{body}\n</REQUESTDATA>\n</IMPORTDATA>\n</BODY>\n</ENVELOPE>\n"
    )


def voucher_xml_message(bill: Bill, party: Party, s: TallySettings, item_units: Optional[dict[str, str]] = None) -> str:
    vch_type, acct_ledger, party_sign = _vch(bill, s)
    all_lines = inventory_lines(bill, s, item_units)
    lines = all_lines if s.inventory else []
    total = round(float(bill.amount), 2)
    # Tally: debit = negative amount + ISDEEMEDPOSITIVE Yes ; credit = positive + No
    if party_sign == "Dr":
        party_amt, party_pos, acct_amt, acct_pos = -total, "Yes", total, "No"
    else:
        party_amt, party_pos, acct_amt, acct_pos = total, "No", -total, "Yes"
    d = _tally_date(bill.bill_date)
    # Invoice mode (stock items): Tally's own export uses <LEDGERENTRIES.LIST> + <ALLINVENTORYENTRIES.LIST>;
    # voucher mode (accounting only) uses <ALLLEDGERENTRIES.LIST>. Mirror that exactly.
    view = "Invoice Voucher View" if lines else "Accounting Voucher View"
    led_tag = "LEDGERENTRIES.LIST" if lines else "ALLLEDGERENTRIES.LIST"
    p = []
    p.append('<TALLYMESSAGE xmlns:UDF="TallyUDF">')
    p.append(f'<VOUCHER VCHTYPE="{_x(vch_type)}" ACTION="Create" OBJVIEW="{view}">')
    p.append(f"<DATE>{d}</DATE><EFFECTIVEDATE>{d}</EFFECTIVEDATE>")
    p.append(f"<VOUCHERTYPENAME>{_x(vch_type)}</VOUCHERTYPENAME>")
    p.append(f"<VOUCHERNUMBER>{_x(bill.bill_no)}</VOUCHERNUMBER><REFERENCE>{_x(bill.bill_no)}</REFERENCE>")
    p.append(f"<PARTYLEDGERNAME>{_x(party.name)}</PARTYLEDGERNAME><PARTYNAME>{_x(party.name)}</PARTYNAME>")
    p.append(f"<BASICBASEPARTYNAME>{_x(party.name)}</BASICBASEPARTYNAME>")
    if party_sign == "Dr":  # buyer name only makes sense on a sales invoice
        p.append(f"<BASICBUYERNAME>{_x(party.name)}</BASICBUYERNAME>")
    p.append(f"<PERSISTEDVIEW>{view}</PERSISTEDVIEW>")
    p.append(f"<VCHENTRYMODE>{'Item Invoice' if lines else 'As Voucher'}</VCHENTRYMODE>")
    p.append(f"<ISINVOICE>{'Yes' if lines else 'No'}</ISINVOICE>")
    p.append(f"<NARRATION>{_x(_narration(bill, party, all_lines, include_items=not lines))}</NARRATION>")
    # party ledger with bill-wise reference
    p.append(f"<{led_tag}>")
    p.append(f"<LEDGERNAME>{_x(party.name)}</LEDGERNAME><ISDEEMEDPOSITIVE>{party_pos}</ISDEEMEDPOSITIVE>")
    p.append(f"<ISPARTYLEDGER>Yes</ISPARTYLEDGER><AMOUNT>{_amt(party_amt)}</AMOUNT>")
    p.append(f"<BILLALLOCATIONS.LIST><NAME>{_x(bill.bill_no)}</NAME><BILLTYPE>New Ref</BILLTYPE>"
             f"<AMOUNT>{_amt(party_amt)}</AMOUNT></BILLALLOCATIONS.LIST>")
    p.append(f"</{led_tag}>")
    if lines:
        for ln in lines:
            amt = ln.amount if acct_pos == "No" else -ln.amount
            p.append("<ALLINVENTORYENTRIES.LIST>")
            p.append(f"<STOCKITEMNAME>{_x(ln.item)}</STOCKITEMNAME><ISDEEMEDPOSITIVE>{acct_pos}</ISDEEMEDPOSITIVE>")
            p.append(f"<RATE>{_amt(ln.rate)}/{_x(ln.unit)}</RATE><AMOUNT>{_amt(amt)}</AMOUNT>")
            p.append(f"<ACTUALQTY>{_fmt_qty(ln.qty)} {_x(ln.unit)}</ACTUALQTY><BILLEDQTY>{_fmt_qty(ln.qty)} {_x(ln.unit)}</BILLEDQTY>")
            if s.godown:
                p.append(f"<BATCHALLOCATIONS.LIST><GODOWNNAME>{_x(s.godown)}</GODOWNNAME><BATCHNAME>Primary Batch</BATCHNAME>"
                         f"<AMOUNT>{_amt(amt)}</AMOUNT><ACTUALQTY>{_fmt_qty(ln.qty)} {_x(ln.unit)}</ACTUALQTY>"
                         f"<BILLEDQTY>{_fmt_qty(ln.qty)} {_x(ln.unit)}</BILLEDQTY></BATCHALLOCATIONS.LIST>")
            p.append(f"<ACCOUNTINGALLOCATIONS.LIST><LEDGERNAME>{_x(acct_ledger)}</LEDGERNAME>"
                     f"<ISDEEMEDPOSITIVE>{acct_pos}</ISDEEMEDPOSITIVE><AMOUNT>{_amt(amt)}</AMOUNT></ACCOUNTINGALLOCATIONS.LIST>")
            p.append("</ALLINVENTORYENTRIES.LIST>")
        # rounding difference between line amounts and the booked total goes to the sales ledger
        diff = round(total - sum(ln.amount for ln in lines), 2)
        if abs(diff) >= 0.01:
            amt = diff if acct_pos == "No" else -diff
            # The flag must follow the sign: when the booked total is *below* the line total the
            # difference is on the other side (a Dr on Sales / a Cr on Purchase), never "No" with a
            # negative amount.
            pos = "Yes" if amt < 0 else "No"
            p.append(f"<LEDGERENTRIES.LIST><LEDGERNAME>{_x(acct_ledger)}</LEDGERNAME>"
                     f"<ISDEEMEDPOSITIVE>{pos}</ISDEEMEDPOSITIVE><AMOUNT>{_amt(amt)}</AMOUNT></LEDGERENTRIES.LIST>")
    else:
        p.append(f"<{led_tag}><LEDGERNAME>{_x(acct_ledger)}</LEDGERNAME>"
                 f"<ISDEEMEDPOSITIVE>{acct_pos}</ISDEEMEDPOSITIVE><AMOUNT>{_amt(acct_amt)}</AMOUNT></{led_tag}>")
    p.append("</VOUCHER>")
    p.append("</TALLYMESSAGE>")
    return "\n".join(p)


def vouchers_xml(bills: Sequence[Bill], parties: dict[int, Party], s: TallySettings,
                 item_units: Optional[dict[str, str]] = None) -> str:
    msgs = [voucher_xml_message(b, parties[b.party_id], s, item_units) for b in bills]
    return _envelope(msgs, "Vouchers", s.company)


def masters_xml(parties: Sequence[Party], items: Sequence[tuple[str, str, Optional[str]]], s: TallySettings) -> str:
    """Ledger + Unit + StockItem masters. items = (name, unit, stock_group)."""
    msgs = []
    for p in parties:
        addr = "".join(f"<ADDRESS>{_x(a.strip())}</ADDRESS>" for a in (p.address or "").splitlines() if a.strip())
        ob = ""
        if p.opening_balance:
            # opening: Dr negative, Cr positive
            v = -abs(p.opening_balance) if (p.opening_drcr or "Dr") == "Dr" else abs(p.opening_balance)
            ob = f"<OPENINGBALANCE>{_amt(v)}</OPENINGBALANCE>"
        msgs.append(
            f'<TALLYMESSAGE xmlns:UDF="TallyUDF"><LEDGER NAME="{_x(p.name)}" ACTION="Create">'
            f"<NAME.LIST><NAME>{_x(p.name)}</NAME></NAME.LIST><PARENT>{_x((p.group_name or '').strip() or 'Sundry Debtors')}</PARENT>"
            f"<ISBILLWISEON>Yes</ISBILLWISEON>"
            + (f"<ADDRESS.LIST TYPE=\"String\">{addr}</ADDRESS.LIST>" if addr else "")
            + (f"<LEDSTATENAME>{_x(p.state)}</LEDSTATENAME>" if p.state else "")
            + (f"<COUNTRYNAME>{_x(p.country)}</COUNTRYNAME>" if p.country else "")
            + (f"<PINCODE>{_x(p.pincode)}</PINCODE>" if p.pincode else "")
            + ob + "</LEDGER></TALLYMESSAGE>"
        )
    units = sorted({(u or s.default_unit).strip() or "Nos" for _, u, _ in items})
    for u in units:
        msgs.append(f'<TALLYMESSAGE xmlns:UDF="TallyUDF"><UNIT NAME="{_x(u)}" ACTION="Create">'
                    f"<NAME>{_x(u)}</NAME><ISSIMPLEUNIT>Yes</ISSIMPLEUNIT><DECIMALPLACES>4</DECIMALPLACES></UNIT></TALLYMESSAGE>")  # matches the 4 dp quantities
    groups = sorted({(g or s.stock_group).strip() for _, _, g in items if (g or s.stock_group).strip()})
    for g in groups:
        msgs.append(f'<TALLYMESSAGE xmlns:UDF="TallyUDF"><STOCKGROUP NAME="{_x(g)}" ACTION="Create">'
                    f"<NAME.LIST><NAME>{_x(g)}</NAME></NAME.LIST><PARENT/></STOCKGROUP></TALLYMESSAGE>")
    for name, unit, grp in items:
        unit = (unit or s.default_unit).strip() or "Nos"  # BASEUNITS is mandatory for a stock item
        grp = (grp or s.stock_group).strip()
        msgs.append(f'<TALLYMESSAGE xmlns:UDF="TallyUDF"><STOCKITEM NAME="{_x(name)}" ACTION="Create">'
                    f"<NAME.LIST><NAME>{_x(name)}</NAME></NAME.LIST>"
                    + (f"<PARENT>{_x(grp)}</PARENT>" if grp else "")
                    + f"<BASEUNITS>{_x(unit)}</BASEUNITS></STOCKITEM></TALLYMESSAGE>")
    return _envelope(msgs, "All Masters", s.company, dups="@@DUPIGNORECOMBINE")


# --------------------------------------------------------------------------- #
# Excel / CSV (one row per entry)
# --------------------------------------------------------------------------- #
VOUCHER_HEADERS = [
    "Voucher Date", "Voucher Type Name", "Voucher Number", "Party Ledger Name",
    "Ledger Name", "Ledger Amount", "Ledger Amount Dr/Cr", "Bill Name", "Bill Type",
    "Stock Item Name", "Quantity", "Unit", "Rate", "Item Amount", "Narration",
]


def voucher_rows(bills: Sequence[Bill], parties: dict[int, Party], s: TallySettings,
                 item_units: Optional[dict[str, str]] = None) -> list[list]:
    rows: list[list] = []
    for b in bills:
        party = parties[b.party_id]
        vch_type, acct_ledger, party_sign = _vch(b, s)
        acct_sign = "Cr" if party_sign == "Dr" else "Dr"
        all_lines = inventory_lines(b, s, item_units)
        lines = all_lines if s.inventory else []
        total = round(float(b.amount), 2)
        narr = _narration(b, party, all_lines, include_items=not lines)
        head = [b.bill_date, vch_type, b.bill_no, party.name]
        rows.append(head + [party.name, total, party_sign, b.bill_no, "New Ref", None, None, None, None, None, narr])
        if lines:
            for ln in lines:
                rows.append(head + [acct_ledger, ln.amount, acct_sign, None, None,
                                    ln.item, ln.qty, ln.unit, round(ln.rate, 2), ln.amount, None])
            diff = round(total - sum(ln.amount for ln in lines), 2)
            if abs(diff) >= 0.01:
                # a negative difference is booked on the opposite side, never as a negative amount
                sign = acct_sign if diff > 0 else ("Dr" if acct_sign == "Cr" else "Cr")
                rows.append(head + [acct_ledger, abs(diff), sign, None, None, None, None, None, None, None, "balancing difference vs booked total"])
        else:
            rows.append(head + [acct_ledger, total, acct_sign, None, None, None, None, None, None, None, None])
    return rows


def vouchers_xlsx(bills, parties, s, item_units=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vouchers"
    for i, h in enumerate(VOUCHER_HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="FFFEC530")
    for r, row in enumerate(voucher_rows(bills, parties, s, item_units), start=2):
        for ci, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=ci, value=v if not hasattr(v, "year") else datetime.combine(v, datetime.min.time()))
            if ci == 1:
                c.number_format = "dd-mm-yyyy"
            elif ci in (6, 13, 14):
                c.number_format = "#,##0.00"
            elif ci == 11:
                c.number_format = "General"  # "0.###" would show a whole number as "13." in Excel
    for i, w in enumerate([12, 16, 14, 26, 22, 14, 9, 12, 9, 28, 10, 7, 11, 13, 50], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    rm = wb.create_sheet("ReadMe")
    for i, line in enumerate([
        "How to import into TallyPrime",
        "1. Gateway of Tally → Import → Transactions → select this file (Excel).",
        "2. Map the columns once (Voucher Type Name, Voucher Date, Voucher Number, Ledger Name, Ledger Amount, Ledger Amount Dr/Cr,",
        "   Bill Name / Bill Type, Stock Item Name, Quantity, Rate, Amount, Narration). Save the mapping template; it is reused next time.",
        "3. One voucher = all rows with the same Voucher Number: first row = party ledger (Dr for sales / Cr for purchase, New Ref bill),",
        "   following rows = sales/purchase ledger allocations, one per stock item.",
        "4. Ledgers, units and stock items must exist in Tally first – import the Masters XML from the same export if they don't.",
        "5. Prefer the Voucher XML file from the same export if your Tally release cannot import Excel vouchers.",
    ], start=1):
        rm.cell(row=i, column=1, value=line)
    rm.column_dimensions["A"].width = 130
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def vouchers_csv(bills, parties, s, item_units=None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(VOUCHER_HEADERS)
    for row in voucher_rows(bills, parties, s, item_units):
        w.writerow([(v.strftime("%d-%m-%Y") if hasattr(v, "year") else ("" if v is None else v)) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


# --------------------------------------------------------------------------- #
# bundle
# --------------------------------------------------------------------------- #
FORMATS = {
    "voucher_xml": "Voucher XML (Tally import)",
    "masters_xml": "Masters XML (ledgers, units, stock items)",
    "voucher_xlsx_official": "Voucher Excel – Tally 'Accounting Voucher' official template",
    "voucher_xlsx": "Voucher Excel – row layout (needs one-time mapping)",
    "voucher_csv": "Voucher CSV",
}


@dataclass
class ExportBundle:
    files: dict[str, bytes] = field(default_factory=dict)  # filename -> bytes
    bill_ids: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def zip_bytes(self) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in self.files.items():
                z.writestr(name, data)
        return buf.getvalue()


def _item_units(db_path=None) -> dict[str, str]:
    units: dict[str, str] = {}
    for it in db.list_items(db_path):
        if it.get("default_unit"):
            units[it["name"].lower()] = it["default_unit"]
            for a in it.get("aliases") or []:
                units.setdefault(a.lower(), it["default_unit"])
    return units


def build_bundle(bills: Sequence[Bill], formats: Sequence[str], s: Optional[TallySettings] = None,
                 db_path=None, stem: Optional[str] = None) -> ExportBundle:
    s = s or TallySettings.load(db_path)
    parties = {p.id: p for p in db.list_parties(db_path)}
    missing = [b for b in bills if b.party_id not in parties]
    bundle = ExportBundle(bill_ids=[b.id for b in bills if b.id is not None and b.party_id in parties])
    if missing:
        bundle.warnings.append(f"{len(missing)} bill(s) skipped – party no longer exists.")
        bills = [b for b in bills if b.party_id in parties]
    if not bills:
        return bundle
    item_units = _item_units(db_path)
    stem = stem or f"Tally_Vouchers_{datetime.now():%Y%m%d_%H%M%S}"
    if "voucher_xml" in formats:
        bundle.files[f"{stem}.xml"] = vouchers_xml(bills, parties, s, item_units).encode("utf-8")
    if "masters_xml" in formats:
        used_parties = {b.party_id for b in bills}
        items: dict[str, tuple[str, str, Optional[str]]] = {}
        item_groups = {i["name"].lower(): i.get("group_name") for i in db.list_items(db_path)}
        if s.inventory:
            for b in bills:
                for ln in inventory_lines(b, s, item_units):
                    items.setdefault(ln.item.lower(), (ln.item, ln.unit, item_groups.get(ln.item.lower())))
        bundle.files[f"{stem}_masters.xml"] = masters_xml(
            [parties[pid] for pid in sorted(used_parties)], list(items.values()), s
        ).encode("utf-8")
    if "voucher_xlsx_official" in formats:
        bundle.files[f"{stem}_TallyTemplate.xlsx"] = official_xlsx(official_rows_for_bills(bills, parties, s, item_units))
    if "voucher_xlsx" in formats:
        bundle.files[f"{stem}.xlsx"] = vouchers_xlsx(bills, parties, s, item_units)
    if "voucher_csv" in formats:
        bundle.files[f"{stem}.csv"] = vouchers_csv(bills, parties, s, item_units)
    if s.inventory:
        seen_units: dict[str, set[str]] = {}
        for b in bills:
            lines = inventory_lines(b, s, item_units)
            if not lines:
                bundle.warnings.append(f"Bill {b.bill_no} ({parties[b.party_id].name}) has no usable line items – exported as an accounting-only voucher.")
            for ln in lines:
                seen_units.setdefault(ln.item.lower(), set()).add(ln.unit)
        bundle.warnings.extend(unit_warnings(seen_units, item_units))
    return bundle


def unit_warnings(seen_units: dict[str, set[str]], item_units: dict[str, str]) -> list[str]:
    """Tally rejects a voucher whose quantity unit is not the stock item's base unit (or one of its
    alternate units). Flag items exported in a unit that differs from the items master, and items
    exported in more than one unit within the same export."""
    out: list[str] = []
    for key, units in sorted(seen_units.items()):
        master = (item_units.get(key) or "").strip()
        exported = sorted(units)
        if len(exported) > 1:
            out.append(f"Item '{key}' is exported in several units ({', '.join(exported)}); a Tally stock item has one "
                       "base unit – the vouchers in the other unit will be rejected on import.")
        if master and any(u.lower() != master.lower() for u in exported):
            out.append(f"Item '{key}': items master unit is '{master}' but it is exported as "
                       f"{' / '.join(u for u in exported if u.lower() != master.lower())}. If the stock item exists in Tally "
                       f"with base unit '{master}', the voucher will be rejected – add '{', '.join(exported)}' as an alternate "
                       "unit in Tally or enter the line in the master unit.")
    return out


def validate_xml(data: bytes) -> int:
    """Parse the XML and return the number of <VOUCHER> elements (raises on malformed XML)."""
    root = ET.fromstring(data)
    return len(root.findall(".//VOUCHER"))


# --------------------------------------------------------------------------- #
# Tally's official "Accounting Voucher" Excel template (Import → Transactions, no mapping needed)
# --------------------------------------------------------------------------- #
OFFICIAL_TEMPLATE = config.FIXTURES_DIR / "tally_templates" / "AccountingVouchers.xlsx"
# The 14 columns of Tally's sample sheet, in its order (the last header carries a trailing space in the
# original file – we keep whatever the template file says). Then optional columns, spelled exactly as in
# the template's Read Me sheet.
OFFICIAL_BASE_HEADERS = [
    "Voucher Date", "Voucher Type Name", "Voucher Number", "Buyer/Supplier - Address", "Buyer/Supplier - Pincode",
    "Ledger Name", "Ledger Amount", "Ledger Amount Dr/Cr", "Item Name", "Billed Quantity", "Item Rate",
    "Item Rate per", "Item Amount", "Change Mode",
]
OFFICIAL_EXTRA_HEADERS = [
    "Quantity UOM", "Bill Type of Ref", "Bill Name", "Bill Amount", "Bill Amount - Dr/Cr", "Reference No.",
    "Voucher Narration", "Bank Allocations - Transaction Type", "Bank Allocations - Amount",
    "Bank Allocations - Inst No.", "Bank Allocations - Inst Date", "Bank Allocations - Bank Name",
    "Bank Allocations - Branch",
]
OFFICIAL_HEADERS = OFFICIAL_BASE_HEADERS + OFFICIAL_EXTRA_HEADERS


def _orow(**kw) -> dict:
    """One row of the official sheet (keys = header names)."""
    row = {h: None for h in OFFICIAL_HEADERS}
    for k, v in kw.items():
        if k not in row:
            raise KeyError(k)
        row[k] = v
    return row


def official_rows_for_bills(bills: Sequence[Bill], parties: dict[int, Party], s: TallySettings,
                            item_units: Optional[dict[str, str]] = None) -> list[dict]:
    rows: list[dict] = []
    for b in bills:
        party = parties[b.party_id]
        vch_type, acct_ledger, party_sign = _vch(b, s)
        acct_sign = "Cr" if party_sign == "Dr" else "Dr"
        all_lines = inventory_lines(b, s, item_units)
        lines = all_lines if s.inventory else []
        total = round(float(b.amount), 2)
        mode = "Item Invoice" if lines else "As Voucher"
        head = {
            "Voucher Date": b.bill_date, "Voucher Type Name": vch_type, "Voucher Number": b.bill_no,
            "Reference No.": b.bill_no, "Change Mode": mode,
        }
        addr = ", ".join(x.strip() for x in (party.address or "").splitlines() if x.strip()) or None
        rows.append(_orow(**head, **{
            "Buyer/Supplier - Address": addr, "Buyer/Supplier - Pincode": party.pincode or None,
            "Ledger Name": party.name, "Ledger Amount": total, "Ledger Amount Dr/Cr": party_sign,
            "Bill Type of Ref": "New Ref", "Bill Name": b.bill_no, "Bill Amount": total, "Bill Amount - Dr/Cr": party_sign,
            "Voucher Narration": _narration(b, party, all_lines, include_items=not lines),
        }))
        if lines:
            for ln in lines:
                rows.append(_orow(**head, **{
                    "Ledger Name": acct_ledger, "Ledger Amount": ln.amount, "Ledger Amount Dr/Cr": acct_sign,
                    "Item Name": ln.item, "Billed Quantity": ln.qty, "Quantity UOM": ln.unit,
                    "Item Rate": round(ln.rate, 2), "Item Rate per": ln.unit, "Item Amount": ln.amount,
                }))
            diff = round(total - sum(ln.amount for ln in lines), 2)
            if abs(diff) >= 0.01:
                sign = acct_sign if diff > 0 else ("Dr" if acct_sign == "Cr" else "Cr")
                rows.append(_orow(**head, **{"Ledger Name": acct_ledger, "Ledger Amount": abs(diff), "Ledger Amount Dr/Cr": sign}))
        else:
            rows.append(_orow(**head, **{"Ledger Name": acct_ledger, "Ledger Amount": total, "Ledger Amount Dr/Cr": acct_sign}))
    return rows


def official_xlsx(rows: list[dict]) -> bytes:
    """Write rows into a copy of the active Accounting Voucher template.

    Official template (default): its own columns + our optional Tally columns are appended.
    Custom template uploaded by the user: ONLY its columns are filled (nothing is added), so the
    file looks exactly like the template they gave us.
    """
    from . import templates

    template_headers: list[Optional[str]] = []  # by column position; None = blank header cell
    tpath = templates.active_path("voucher")
    custom = templates.is_custom("voucher")
    if tpath is not None:
        wb = openpyxl.load_workbook(tpath)  # read fully into memory; the file handle is closed by openpyxl
        ws = wb["Accounting Voucher"] if "Accounting Voucher" in wb.sheetnames else wb.worksheets[0]
        template_headers = [None if c.value is None or not str(c.value).strip() else str(c.value) for c in ws[1]]
        while template_headers and template_headers[-1] is None:
            template_headers.pop()
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    else:  # fallback: bare workbook with the same headers
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Accounting Voucher"
    # header row = template's own cells (exact spelling incl. trailing spaces, blank columns kept in
    # place so every value lands under its own header) + our optional columns
    headers: list[Optional[str]] = template_headers[:] if template_headers else OFFICIAL_BASE_HEADERS[:]
    norm = {h.strip(): i for i, h in enumerate(headers) if h is not None}
    if not custom:
        for h in OFFICIAL_HEADERS:
            if h not in norm:
                headers.append(h)
                norm[h] = len(headers) - 1
    bold = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        if c.value is None and h is not None:
            c.value = h
            c.font = bold if h in ("Voucher Date", "Voucher Type Name", "Ledger Name", "Ledger Amount") else Font()
    for r, row in enumerate(rows, start=2):
        for h, v in row.items():
            if h not in norm:  # custom template without this column -> value is left out
                continue
            ci = norm[h] + 1
            if v is None:
                continue
            if hasattr(v, "year"):
                c = ws.cell(row=r, column=ci, value=datetime.combine(v, datetime.min.time()))
                c.number_format = "dd-mm-yyyy"
            else:
                c = ws.cell(row=r, column=ci, value=v)
                if h in ("Ledger Amount", "Item Amount", "Bill Amount", "Bank Allocations - Amount", "Item Rate"):
                    c.number_format = "#,##0.00"
    for i, h in enumerate(headers, start=1):
        if h is None:
            continue
        cur = ws.column_dimensions[get_column_letter(i)].width
        if not cur or cur < 12:
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 2))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Simple (non-inventory) vouchers – Contra / Receipt / Payment from bank slips
# --------------------------------------------------------------------------- #
@dataclass
class BankAlloc:
    txn_type: str = "Cash"  # Cash, Cheque/DD, e-Fund Transfer, ...
    inst_no: str = ""
    inst_date: Optional[object] = None  # date
    bank_name: str = ""
    branch: str = ""


@dataclass
class SimpleEntry:
    ledger: str
    amount: float  # positive
    drcr: str  # "Dr" / "Cr"
    is_party: bool = False
    bill_type: str = ""  # "On Account" / "Agst Ref" / "New Ref" – party entries only
    bill_name: str = ""
    bank: Optional[BankAlloc] = None  # bank ledger entries only


@dataclass
class SimpleVoucher:
    date: object  # date
    vch_type: str
    number: str
    entries: list[SimpleEntry]
    narration: str = ""
    reference: str = ""

    def balanced(self) -> bool:
        dr = sum(e.amount for e in self.entries if e.drcr == "Dr")
        cr = sum(e.amount for e in self.entries if e.drcr == "Cr")
        return abs(dr - cr) < 0.005


def simple_voucher_xml_message(v: SimpleVoucher) -> str:
    d = _tally_date(v.date)
    p = ['<TALLYMESSAGE xmlns:UDF="TallyUDF">',
         f'<VOUCHER VCHTYPE="{_x(v.vch_type)}" ACTION="Create" OBJVIEW="Accounting Voucher View">',
         f"<DATE>{d}</DATE><EFFECTIVEDATE>{d}</EFFECTIVEDATE>",
         f"<VOUCHERTYPENAME>{_x(v.vch_type)}</VOUCHERTYPENAME>",
         f"<VOUCHERNUMBER>{_x(v.number)}</VOUCHERNUMBER>"]
    if v.reference:
        p.append(f"<REFERENCE>{_x(v.reference)}</REFERENCE>")
    party = next((e for e in v.entries if e.is_party), None)
    if party:
        p.append(f"<PARTYLEDGERNAME>{_x(party.ledger)}</PARTYLEDGERNAME><PARTYNAME>{_x(party.ledger)}</PARTYNAME>")
    p.append("<PERSISTEDVIEW>Accounting Voucher View</PERSISTEDVIEW><VCHENTRYMODE>As Voucher</VCHENTRYMODE><ISINVOICE>No</ISINVOICE>")
    if v.narration:
        p.append(f"<NARRATION>{_x(v.narration)}</NARRATION>")
    for e in v.entries:
        amt = -abs(e.amount) if e.drcr == "Dr" else abs(e.amount)
        pos = "Yes" if e.drcr == "Dr" else "No"
        p.append("<ALLLEDGERENTRIES.LIST>")
        p.append(f"<LEDGERNAME>{_x(e.ledger)}</LEDGERNAME><ISDEEMEDPOSITIVE>{pos}</ISDEEMEDPOSITIVE>")
        if e.is_party:
            p.append("<ISPARTYLEDGER>Yes</ISPARTYLEDGER>")
        p.append(f"<AMOUNT>{_amt(amt)}</AMOUNT>")
        if e.is_party and e.bill_type:
            # Bill-wise party ledgers need an allocation; Tally's own export writes an "On Account"
            # allocation without a NAME, and a NAME only for New Ref / Agst Ref / Advance.
            name = "" if e.bill_type == "On Account" else f"<NAME>{_x(e.bill_name or v.number)}</NAME>"
            p.append(f"<BILLALLOCATIONS.LIST>{name}<BILLTYPE>{_x(e.bill_type)}</BILLTYPE>"
                     f"<AMOUNT>{_amt(amt)}</AMOUNT></BILLALLOCATIONS.LIST>")
        if e.bank is not None:
            b = e.bank
            p.append("<BANKALLOCATIONS.LIST>")
            p.append(f"<DATE>{d}</DATE>")
            if b.inst_date is not None:
                p.append(f"<INSTRUMENTDATE>{_tally_date(b.inst_date)}</INSTRUMENTDATE>")
            p.append(f"<TRANSACTIONTYPE>{_x(b.txn_type or 'Cash')}</TRANSACTIONTYPE>")
            if b.bank_name:
                p.append(f"<BANKNAME>{_x(b.bank_name)}</BANKNAME>")
            if b.inst_no:
                p.append(f"<INSTRUMENTNUMBER>{_x(b.inst_no)}</INSTRUMENTNUMBER>")
            p.append(f"<AMOUNT>{_amt(amt)}</AMOUNT></BANKALLOCATIONS.LIST>")
        p.append("</ALLLEDGERENTRIES.LIST>")
    p.append("</VOUCHER></TALLYMESSAGE>")
    return "\n".join(p)


def simple_vouchers_xml(vouchers: Sequence[SimpleVoucher], s: TallySettings) -> str:
    return _envelope([simple_voucher_xml_message(v) for v in vouchers], "Vouchers", s.company)


def official_rows_for_simple(vouchers: Sequence[SimpleVoucher]) -> list[dict]:
    rows: list[dict] = []
    for v in vouchers:
        head = {"Voucher Date": v.date, "Voucher Type Name": v.vch_type, "Voucher Number": v.number,
                "Change Mode": "As Voucher", "Reference No.": v.reference or None}
        for i, e in enumerate(v.entries):
            extra = {"Ledger Name": e.ledger, "Ledger Amount": round(abs(e.amount), 2), "Ledger Amount Dr/Cr": e.drcr}
            if i == 0:
                extra["Voucher Narration"] = v.narration or None
            if e.is_party and e.bill_type:
                extra.update({"Bill Type of Ref": e.bill_type, "Bill Name": e.bill_name or (v.number if e.bill_type != "On Account" else None),
                              "Bill Amount": round(abs(e.amount), 2), "Bill Amount - Dr/Cr": e.drcr})
            if e.bank is not None:
                extra.update({"Bank Allocations - Transaction Type": e.bank.txn_type or "Cash",
                              "Bank Allocations - Amount": round(abs(e.amount), 2),
                              "Bank Allocations - Inst No.": e.bank.inst_no or None,
                              "Bank Allocations - Inst Date": e.bank.inst_date,
                              "Bank Allocations - Bank Name": e.bank.bank_name or None,
                              "Bank Allocations - Branch": e.bank.branch or None})
            rows.append(_orow(**head, **extra))
    return rows


def simple_rows(vouchers: Sequence[SimpleVoucher]) -> list[list]:
    """Rows in the app's own layout (VOUCHER_HEADERS) for CSV / row-layout Excel."""
    rows: list[list] = []
    for v in vouchers:
        party = next((e.ledger for e in v.entries if e.is_party), "")
        head = [v.date, v.vch_type, v.number, party]
        for i, e in enumerate(v.entries):
            rows.append(head + [e.ledger, round(abs(e.amount), 2), e.drcr,
                                (e.bill_name or (v.number if e.bill_type not in ("", "On Account") else None)) if e.is_party else None,
                                e.bill_type or None if e.is_party else None,
                                None, None, None, None, None, v.narration if i == 0 else None])
    return rows


def simple_vouchers_xlsx(vouchers: Sequence[SimpleVoucher]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vouchers"
    for i, h in enumerate(VOUCHER_HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="FFFEC530")
    for r, row in enumerate(simple_rows(vouchers), start=2):
        for ci, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=ci, value=v if not hasattr(v, "year") else datetime.combine(v, datetime.min.time()))
            if ci == 1:
                c.number_format = "dd-mm-yyyy"
            elif ci == 6:
                c.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def simple_vouchers_csv(vouchers: Sequence[SimpleVoucher]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(VOUCHER_HEADERS)
    for row in simple_rows(vouchers):
        w.writerow([(v.strftime("%d-%m-%Y") if hasattr(v, "year") else ("" if v is None else v)) for v in row])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def bank_ledger_masters_xml(bank_ledgers: Sequence[str], parties: Sequence[Party], s: TallySettings) -> str:
    """Masters for bank vouchers: bank ledgers (group Bank Accounts) + any party ledgers used."""
    msgs = []
    for name in sorted({b.strip() for b in bank_ledgers if b and b.strip()}):
        msgs.append(f'<TALLYMESSAGE xmlns:UDF="TallyUDF"><LEDGER NAME="{_x(name)}" ACTION="Create">'
                    f"<NAME.LIST><NAME>{_x(name)}</NAME></NAME.LIST><PARENT>Bank Accounts</PARENT></LEDGER></TALLYMESSAGE>")
    inner = masters_xml(parties, [], s) if parties else ""
    if inner:
        start = inner.index("<REQUESTDATA>") + len("<REQUESTDATA>")
        end = inner.index("</REQUESTDATA>")
        msgs.append(inner[start:end].strip())
    return _envelope(msgs, "All Masters", s.company, dups="@@DUPIGNORECOMBINE")


def build_simple_bundle(vouchers: Sequence[SimpleVoucher], formats: Sequence[str], s: TallySettings,
                        stem: str, bank_ledgers: Sequence[str] = (), parties: Sequence[Party] = ()) -> ExportBundle:
    bundle = ExportBundle()
    if not vouchers:
        return bundle
    for v in vouchers:
        if not v.balanced():
            bundle.warnings.append(f"Voucher {v.number} does not balance – check its amounts.")
    if "voucher_xml" in formats:
        bundle.files[f"{stem}.xml"] = simple_vouchers_xml(vouchers, s).encode("utf-8")
    if "masters_xml" in formats and (bank_ledgers or parties):
        bundle.files[f"{stem}_masters.xml"] = bank_ledger_masters_xml(bank_ledgers, parties, s).encode("utf-8")
    if "voucher_xlsx_official" in formats:
        bundle.files[f"{stem}_TallyTemplate.xlsx"] = official_xlsx(official_rows_for_simple(vouchers))
    if "voucher_xlsx" in formats:
        bundle.files[f"{stem}.xlsx"] = simple_vouchers_xlsx(vouchers)
    if "voucher_csv" in formats:
        bundle.files[f"{stem}.csv"] = simple_vouchers_csv(vouchers)
    return bundle
