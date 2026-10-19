"""Write bills to an Excel workbook in Tally Prime's 'Ledgers in Default Template' layout.

Sheet 'Ledger' reproduces the template exactly (we start from the template file so
header styling / column widths / number formats are preserved). Sheet 'Items' carries
the line items for the business's own records.
"""
from __future__ import annotations

import io
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Optional, Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, db
from .importer import HEADERS
from .models import Bill, Party

DATE_FMT = "dd/mmm/yyyy"
AMT_FMT = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '

ITEM_HEADERS = [
    "Party", "Bill No", "Bill Date", "Txn", "Item", "Qty", "Unit", "Weight (Qtl)", "Rate", "Amount",
    "Struck Out", "Source", "Bill Total",
]


def _blank_template_workbook() -> openpyxl.Workbook:
    from . import templates

    tpath = templates.custom_path("masters") if templates.is_custom("masters") else config.TEMPLATE_PATH
    if tpath.exists():
        wb = openpyxl.load_workbook(tpath)
        ws = wb["Ledger"] if "Ledger" in wb.sheetnames else wb.worksheets[0]
        ws.title = "Ledger"
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        # remove any stray sheets
        for extra in [s for s in wb.sheetnames if s != "Ledger"]:
            del wb[extra]
        return wb
    # fallback: build header from scratch
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger"
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFFEC530")
    return wb


def _party_order_key(p: Party, bills: list[Bill]):
    first = min((b.bill_date for b in bills), default=None)
    return (first or datetime.max.date(), p.name.lower())


def build_workbook(
    bills: Sequence[Bill],
    parties: dict[int, Party],
    include_opening_balance: bool = True,
    include_address: bool = True,
    include_items_sheet: bool = True,
    parties_without_bills: Sequence[Party] = (),
) -> openpyxl.Workbook:
    wb = _blank_template_workbook()
    ws = wb["Ledger"]
    # Columns are addressed by header name so a user-supplied template with a different column
    # order (or extra columns) still gets every field in the right place.
    col = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None and str(c.value).strip()}
    if not col:  # header row missing altogether -> Tally's default positions
        col = {h: i for i, h in enumerate(HEADERS, start=1)}
    # A header the template does not have is left out (never written into another column's position).

    def put(r: int, header: str, value):
        return ws.cell(row=r, column=col[header], value=value) if header in col else None

    by_party: dict[int, list[Bill]] = {}
    for b in bills:
        by_party.setdefault(b.party_id, []).append(b)
    for pid in by_party:
        by_party[pid].sort(key=lambda b: (b.bill_date, b.id or 0))

    blocks: list[tuple[Party, list[Bill]]] = []
    for pid, pb in by_party.items():
        p = parties.get(pid)
        if p is None:
            raise ValueError(f"party id {pid} missing from parties map")
        blocks.append((p, pb))
    blocks.sort(key=lambda t: _party_order_key(*t))
    for p in parties_without_bills:
        if p.id not in by_party:
            blocks.append((p, []))

    row = 2
    for p, pb in blocks:
        addr_lines = [ln.strip() for ln in (p.address or "").splitlines() if ln.strip()] if include_address else []
        n_rows = max(1, len(pb), len(addr_lines))
        # first row: ledger master
        put(row, "Name", p.name)
        put(row, "Group Name", p.group_name)
        if include_opening_balance and p.opening_balance is not None:
            c = put(row, "Ledger - Opening Balance", float(p.opening_balance))
            if c is not None:
                c.number_format = AMT_FMT
            put(row, "Ledger Opening Balance - Dr/Cr", p.opening_drcr or ("Cr" if "creditor" in p.group_name.lower() else "Dr"))
        if include_address:
            if p.state:
                put(row, "State", p.state)
            if p.country:
                put(row, "Country", p.country)
            if p.pincode:
                put(row, "Pincode", str(p.pincode))
        for i in range(n_rows):
            r = row + i
            if i < len(pb):
                b = pb[i]
                dc = put(r, "Bill - Date", datetime.combine(b.bill_date, datetime.min.time()))
                if dc is not None:
                    dc.number_format = DATE_FMT
                put(r, "Bill - Name", b.bill_no)
                ac = put(r, "Bill - Amount", round(float(b.amount), 2))
                if ac is not None:
                    ac.number_format = AMT_FMT
                put(r, "Bill Amount - Dr/Cr", b.drcr)
            if i < len(addr_lines):
                put(r, "Address", addr_lines[i])
        row += n_rows

    if include_items_sheet:
        _write_items_sheet(wb, blocks)
    return wb


def _write_items_sheet(wb: openpyxl.Workbook, blocks: list[tuple[Party, list[Bill]]]) -> None:
    ws = wb.create_sheet("Items")
    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="FFFEC530")
    for i, h in enumerate(ITEM_HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = hdr_font, hdr_fill
    r = 2
    for p, pb in blocks:
        for b in pb:
            if not b.items:
                continue
            for it in b.items:
                vals = [
                    p.name, b.bill_no, datetime.combine(b.bill_date, datetime.min.time()), b.txn_type,
                    it.item, it.qty, it.unit, it.weight_qtl, it.rate, it.effective_amount,
                    "Yes" if it.struck_out else "", b.source, round(float(b.amount), 2),
                ]
                for ci, v in enumerate(vals, start=1):
                    c = ws.cell(row=r, column=ci, value=v)
                    if ci == 3:
                        c.number_format = DATE_FMT
                    elif ci in (9, 10, 13):
                        c.number_format = "#,##0.00"
                    elif ci in (6, 8):
                        c.number_format = "0.00"
                if it.struck_out:
                    for ci in range(1, len(vals) + 1):
                        ws.cell(row=r, column=ci).font = Font(strike=True, color="FF888888")
                r += 1
    widths = [24, 12, 13, 9, 26, 8, 7, 12, 11, 13, 10, 8, 13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def export_bills(
    bill_ids: Optional[Sequence[int]] = None,
    db_path=None,
    date_from=None,
    date_to=None,
    only_unexported: bool = False,
    include_opening_balance: bool = True,
    include_address: bool = True,
    include_items_sheet: bool = True,
    mark_exported: bool = False,
    out_path: Optional[Path] = None,
) -> tuple[bytes, list[int]]:
    """Build the workbook. Returns (xlsx bytes, exported bill ids)."""
    bills = db.list_bills(
        db_path,
        date_from=date_from,
        date_to=date_to,
        exported=False if only_unexported else None,
        with_items=include_items_sheet,
    )
    if bill_ids is not None:
        wanted = set(bill_ids)
        bills = [b for b in bills if b.id in wanted]
    parties = {p.id: p for p in db.list_parties(db_path)}
    wb = build_workbook(
        bills,
        parties,
        include_opening_balance=include_opening_balance,
        include_address=include_address,
        include_items_sheet=include_items_sheet,
    )
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if out_path:
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        Path(out_path).write_bytes(data)
    ids = [b.id for b in bills if b.id is not None]
    if mark_exported and ids:
        db.mark_exported(ids, True, db_path)
    return data, ids


def default_export_name(prefix: str = "Tally_Ledgers") -> str:
    return f"{prefix}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
