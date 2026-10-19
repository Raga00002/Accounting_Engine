"""Excel import templates: Tally's official sample workbooks by default, optionally overridden by
templates the user uploads (e.g. a copy exported from their own TallyPrime with extra columns).

kind "voucher" -> Accounting Voucher workbook (Import → Transactions)
kind "masters" -> Ledger master workbook   (Import → Masters → Ledgers)
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import openpyxl

from . import config, tenant


def template_dir() -> Path:
    """Folder for the active tenant's uploaded templates."""
    return tenant.current().template_dir

KINDS = {
    "voucher": {
        "label": "Accounting Voucher (transactions)",
        "official": config.FIXTURES_DIR / "tally_templates" / "AccountingVouchers.xlsx",
        "sheet": "Accounting Voucher",
        "required": ["Voucher Date", "Voucher Type Name", "Ledger Name", "Ledger Amount"],
        "recommended": ["Ledger Amount Dr/Cr", "Voucher Number"],
        "custom_name": "AccountingVouchers.xlsx",
    },
    "masters": {
        "label": "Ledger masters",
        "official": config.FIXTURES_DIR / "tally_templates" / "AllAccountingMasters.xlsx",
        "sheet": "Ledger",
        "required": ["Name", "Group Name"],
        "recommended": ["Bill - Date", "Bill - Name", "Bill - Amount", "Bill Amount - Dr/Cr"],
        "custom_name": "AllAccountingMasters.xlsx",
    },
}


@dataclass
class TemplateInfo:
    kind: str
    source: str  # "custom" | "official" | "built-in"
    path: Optional[Path]
    sheet: str
    headers: list[str]
    warnings: list[str]


def custom_path(kind: str) -> Path:
    return template_dir() / KINDS[kind]["custom_name"]


def active_path(kind: str) -> Optional[Path]:
    k = KINDS[kind]
    if custom_path(kind).exists():
        return custom_path(kind)
    if k["official"].exists():
        return k["official"]
    return None


def is_custom(kind: str) -> bool:
    return custom_path(kind).exists()


def _headers(path: Path, sheet: str) -> tuple[str, list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return ws.title, [str(v).strip() for v in first if v is not None and str(v).strip()]
    finally:
        wb.close()


def _check(kind: str, headers: list[str]) -> list[str]:
    k = KINDS[kind]
    warn = []
    missing = [h for h in k["required"] if h not in headers]
    if missing:
        warn.append(f"Mandatory Tally column(s) missing: {', '.join(missing)} – Tally will reject the file.")
    rec = [h for h in k["recommended"] if h not in headers]
    if rec:
        warn.append(f"Column(s) not in this template will be left out of the export: {', '.join(rec)}.")
    return warn


def info(kind: str) -> TemplateInfo:
    k = KINDS[kind]
    p = active_path(kind)
    if p is None:
        return TemplateInfo(kind, "built-in", None, k["sheet"], [], ["No template file found – a plain workbook with Tally's headers is generated."])
    try:
        sheet, headers = _headers(p, k["sheet"])
    except Exception as e:  # noqa: BLE001
        return TemplateInfo(kind, "custom" if p == custom_path(kind) else "official", p, k["sheet"], [], [f"Template unreadable: {e}"])
    return TemplateInfo(kind, "custom" if p == custom_path(kind) else "official", p, sheet, headers,
                        _check(kind, headers) if p == custom_path(kind) else [])


def set_custom(kind: str, data: bytes) -> TemplateInfo:
    """Validate and store an uploaded template. Raises ValueError when it cannot be used."""
    k = KINDS[kind]
    if not data or data[:2] != b"PK":
        raise ValueError("That is not an .xlsx workbook.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"Could not open the workbook: {e}") from e
    try:
        ws = wb[k["sheet"]] if k["sheet"] in wb.sheetnames else wb.worksheets[0]
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(v).strip() for v in first if v is not None and str(v).strip()]
    finally:
        wb.close()
    if not headers:
        raise ValueError("The first row of the sheet is empty – it must contain Tally's column headers.")
    missing = [h for h in k["required"] if h not in headers]
    if missing:
        raise ValueError(f"Template is missing mandatory column(s): {', '.join(missing)}.")
    template_dir().mkdir(parents=True, exist_ok=True)
    custom_path(kind).write_bytes(data)
    return info(kind)


def clear_custom(kind: str) -> None:
    p = custom_path(kind)
    if p.exists():
        p.unlink()
